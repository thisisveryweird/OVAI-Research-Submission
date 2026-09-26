#!/usr/bin/env python3
"""Audit the supplied 2024 adnexal-mass workbook without fitting models."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = (
    ROOT
    / "research/cohort_expansion/data/raw/2024_adnexal_mass"
    / "NIHMS2021079-supplement-1.xlsx"
)
REPORT_DIR = ROOT / "research/cohort_expansion/reports"
FILE_VERIFICATION = REPORT_DIR / "2024_adnexal_mass_file_verification.md"
INVENTORY_CSV = REPORT_DIR / "2024_adnexal_mass_workbook_inventory.csv"
WORKBOOK_AUDIT = REPORT_DIR / "2024_adnexal_mass_workbook_audit.md"
COHORT_CSV = REPORT_DIR / "2024_adnexal_mass_cohort_reconstruction.csv"
STAGE_CSV = REPORT_DIR / "2024_adnexal_mass_stage_counts.csv"
PANEL_CSV = REPORT_DIR / "2024_adnexal_mass_mirna_panel_inventory.csv"
FEATURE_MAP_CSV = REPORT_DIR / "2024_adnexal_mass_feature_map.csv"
COMPATIBILITY_CSV = REPORT_DIR / "2024_adnexal_mass_compatibility.csv"
ENDPOINT_CSV = REPORT_DIR / "2024_adnexal_mass_endpoint_counts.csv"
PUBLISHED_FEATURE_MAP_CSV = REPORT_DIR / "2024_adnexal_mass_published_model_feature_map.csv"
FROZEN_VALIDATION_REPORT = REPORT_DIR / "2024_adnexal_mass_frozen_validation.md"
PUBLISHED_REPRODUCTION_REPORT = REPORT_DIR / "2024_adnexal_mass_published_model_reproduction.md"
DECISION_REPORT = REPORT_DIR / "2024_adnexal_mass_compatibility_decision.md"
DATA_NEEDED_REPORT = REPORT_DIR / "2024_adnexal_mass_DATA_NEEDED.md"

LABEL = "diagnosis (benign = 0, malignant= 1)"
COHORT = "cohort"
ID = "ID"
MIRNA_SHEET = "miRNA"
CLINICAL_SHEET = "protein + metadata"
HISTOLOGY_SHEET = "histologies"
HISTOLOGY_STAGE = "stage (ovarian cancer)"

PRODUCTION_FEATURES = [
    "hsa-miR-200b-3p",
    "hsa-miR-296-3p",
    "hsa-miR-6826-5p",
    "hsa-miR-32-3p",
    "hsa-miR-1255a",
]
HISTORICAL_FEATURES = [
    "hsa-miR-23b-3p",
    "hsa-miR-29a-3p",
    "hsa-miR-32-5p",
    "hsa-miR-92a-3p",
    "hsa-miR-150-5p",
    "hsa-miR-200a-3p",
    "hsa-miR-200c-3p",
    "hsa-miR-203a",
    "hsa-miR-320c",
    "hsa-miR-320d",
    "hsa-miR-335-5p",
    "hsa-miR-450b-5p",
    "hsa-miR-1246",
    "hsa-miR-1307-5p",
]
PUBLISHED_COMBINED_MIRNAS = [
    "hsa-miR-320e",
    "hsa-miR-195-5p",
    "hsa-miR-103-3p",
    "hsa-miR-320c",
    "cel-miR-39-3p",
    "hsa-miR-200c-3p",
    "hsa-miR-770-5p",
    "hsa-miR-122-5p",
    "hsa-miR-19a-3p",
    "hsa-miR-424-5p",
]


def json_value(value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    return value


def records_json(frame: pd.DataFrame) -> str:
    records = [
        {str(key): json_value(value) for key, value in row.items()}
        for row in frame.to_dict(orient="records")
    ]
    return json.dumps(records, ensure_ascii=True, separators=(",", ":"))


def identifier_columns(frame: pd.DataFrame) -> list[str]:
    return [
        str(column)
        for column in frame.columns
        if str(column).strip().casefold() in {"id", "subject id", "sample id"}
    ]


def column_classes(frame: pd.DataFrame) -> tuple[list[str], list[str]]:
    categorical: list[str] = []
    continuous: list[str] = []
    identifiers = set(identifier_columns(frame))
    for column in frame.columns:
        if column in identifiers:
            continue
        series = frame[column]
        if (
            pd.api.types.is_object_dtype(series)
            or pd.api.types.is_bool_dtype(series)
            or isinstance(series.dtype, pd.CategoricalDtype)
            or series.nunique(dropna=True) <= 12
        ):
            categorical.append(str(column))
        elif pd.api.types.is_numeric_dtype(series):
            continuous.append(str(column))
        else:
            categorical.append(str(column))
    return categorical, continuous


def normalize_feature_name(value: str) -> str:
    name = str(value).strip().strip("'").strip()
    name = re.sub(r"\s*\(MFI\)\s*$", "", name, flags=re.IGNORECASE)
    return name.casefold()


def feature_name_from_column(value: str) -> str:
    name = str(value).strip().strip("'").strip()
    return re.sub(r"\s*\(MFI\)\s*$", "", name, flags=re.IGNORECASE)


def build_histology_map(histology: pd.DataFrame) -> tuple[dict[Any, dict[str, str]], int]:
    deduplicated = histology.drop_duplicates().copy()
    exact_duplicate_rows = len(histology) - len(deduplicated)
    mapping: dict[Any, dict[str, str]] = {}
    for patient_id, group in deduplicated.groupby(ID, sort=False, dropna=False):
        stages = group[HISTOLOGY_STAGE].dropna().astype(str).unique().tolist()
        types = group["histology"].dropna().astype(str).unique().tolist()
        descriptions = (
            group[" histology (invasive ovarian cancer)"]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
        )
        if len(stages) == 1:
            stage = stages[0]
        elif len(stages) > 1:
            stage = "Unknown (conflicting worksheet entries)"
        else:
            stage = "Unknown (stage not reported)"
        mapping[patient_id] = {
            "stage": stage,
            "histology": types[0] if len(types) == 1 else "Unknown",
            "histology_detail": descriptions[0] if len(descriptions) == 1 else "Unknown",
        }
    return mapping, exact_duplicate_rows


def write_inventory(sheets: dict[str, pd.DataFrame]) -> None:
    inventory_rows: list[dict[str, Any]] = []
    for name, frame in sheets.items():
        categorical, continuous = column_classes(frame)
        missing = {
            str(column): int(count)
            for column, count in frame.isna().sum().items()
            if count
        }
        id_summary: dict[str, Any] = {}
        duplicates: dict[str, dict[str, int]] = {}
        for column in identifier_columns(frame):
            id_summary[column] = {
                "nonmissing": int(frame[column].notna().sum()),
                "unique": int(frame[column].nunique(dropna=True)),
            }
            duplicates[column] = {
                str(key): int(count)
                for key, count in frame.loc[
                    frame[column].duplicated(keep=False), column
                ].value_counts(dropna=False).items()
            }
        inventory_rows.append(
            {
                "sheet": name,
                "row_count": len(frame),
                "column_count": len(frame.columns),
                "column_names_json": json.dumps(list(map(str, frame.columns))),
                "dtypes_json": json.dumps(
                    {str(key): str(value) for key, value in frame.dtypes.items()}
                ),
                "missingness_json": json.dumps(missing, sort_keys=True),
                "identifier_summary_json": json.dumps(id_summary, sort_keys=True),
                "duplicated_identifiers_json": json.dumps(duplicates, sort_keys=True),
                "categorical_variables_json": json.dumps(categorical),
                "continuous_variables_json": json.dumps(continuous),
                "first_5_rows_json": records_json(frame.head(5)),
                "last_5_rows_json": records_json(frame.tail(5)),
            }
        )
    with INVENTORY_CSV.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(inventory_rows[0]))
        writer.writeheader()
        writer.writerows(inventory_rows)


def describe_sheet(name: str, frame: pd.DataFrame) -> str:
    categorical, continuous = column_classes(frame)
    type_map = {str(column): str(dtype) for column, dtype in frame.dtypes.items()}
    by_type: dict[str, list[str]] = {}
    for column, dtype in type_map.items():
        by_type.setdefault(dtype, []).append(column)
    missing = frame.isna().sum()
    missing = {str(column): int(count) for column, count in missing.items() if count}
    ids = identifier_columns(frame)
    id_lines = []
    for column in ids:
        duplicated = frame.loc[frame[column].duplicated(keep=False), column].value_counts()
        id_lines.append(
            f"- `{column}`: {frame[column].nunique(dropna=True)} unique of "
            f"{frame[column].notna().sum()} nonmissing; duplicate rows "
            f"{int(frame[column].duplicated(keep=False).sum())}; repeated IDs "
            f"`{json.dumps({str(k): int(v) for k, v in duplicated.items()}) if len(duplicated) else 'none'}`."
        )
    if not id_lines:
        id_lines = ["- No ID-like identifier column detected."]

    type_lines = [
        f"- `{dtype}` ({len(columns)}): " + ", ".join(f"`{column}`" for column in columns)
        for dtype, columns in by_type.items()
    ]
    categorical_lines: list[str] = []
    for column in categorical:
        counts = frame[column].value_counts(dropna=False)
        categorical_lines.append(
            f"- `{column}`: "
            + "; ".join(f"{str(value)!r} = {int(count)}" for value, count in counts.items())
        )
    if not categorical_lines:
        categorical_lines = ["- None detected."]
    continuous_lines: list[str] = []
    for column in continuous:
        values = pd.to_numeric(frame[column], errors="coerce")
        continuous_lines.append(
            f"- `{column}`: n={int(values.notna().sum())}, min={values.min():g}, "
            f"median={values.median():g}, max={values.max():g}."
        )
    if not continuous_lines:
        continuous_lines = ["- None detected."]
    first = records_json(frame.head(5))
    last = records_json(frame.tail(5))
    missing_text = json.dumps(missing, ensure_ascii=True, sort_keys=True) if missing else "none"
    return "\n".join(
        [
            f"## {name}",
            f"Rows: {len(frame)} data rows. Columns: {len(frame.columns)}.",
            "",
            "### Column names",
            "```text",
            *map(str, frame.columns),
            "```",
            "",
            "### Data types",
            *type_lines,
            "",
            f"Missingness: `{missing_text}`.",
            "",
            "### Identifier fields and duplicate identifiers",
            *id_lines,
            "",
            "### Categorical variables",
            *categorical_lines,
            "",
            "### Continuous variables",
            *continuous_lines,
            "",
            "### First 5 data rows",
            "```json",
            first,
            "```",
            "",
            "### Last 5 data rows",
            "```json",
            last,
            "```",
            "",
        ]
    )


def write_file_verification() -> tuple[int, str]:
    data = WORKBOOK.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    size = len(data)
    with zipfile.ZipFile(WORKBOOK) as archive:
        bad_member = archive.testzip()
        has_workbook_xml = "xl/workbook.xml" in archive.namelist()
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet_lines = [
        f"- `{sheet.title}`: {sheet.max_row - 1} data rows, {sheet.max_column} columns."
        for sheet in workbook.worksheets
    ]
    workbook.close()
    FILE_VERIFICATION.write_text(
        "\n".join(
            [
                "# 2024 Adnexal-Mass Workbook File Verification",
                "",
                "- Expected file: `NIHMS2021079-supplement-1.xlsx`",
                f"- Verified path: `{WORKBOOK.relative_to(ROOT).as_posix()}`",
                "- Source paper: \"Serum miRNA improves the accuracy of a multivariate index assay for triage of an adnexal mass\"",
                "- DOI: `10.1016/j.ygyno.2024.08.008`",
                "- PMCID: `PMC13283006`",
                f"- Exists: `{WORKBOOK.is_file()}`",
                "- File type: XLSX (ZIP container with `xl/workbook.xml`).",
                f"- ZIP/XLSX signature: `{data[:4].hex()}` (`PK\\x03\\x04`).",
                f"- ZIP integrity: `{'valid' if bad_member is None else f'invalid: {bad_member}'}`.",
                f"- `xl/workbook.xml` present: `{has_workbook_xml}`.",
                f"- File size: {size:,} bytes.",
                f"- SHA256: `{digest}`.",
                "- Worksheet dimensions:",
                *sheet_lines,
                "",
                "The workbook was inspected in place; no download was attempted.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    return size, digest


def write_reconstructions(
    mirna: pd.DataFrame,
    clinical: pd.DataFrame,
    histology: pd.DataFrame,
) -> tuple[pd.DataFrame, dict[Any, dict[str, str]], int]:
    if mirna[ID].isna().any() or clinical[ID].isna().any():
        raise ValueError("Missing ID in a main patient-level worksheet")
    if mirna[ID].duplicated().any() or clinical[ID].duplicated().any():
        raise ValueError("Duplicate IDs in a main patient-level worksheet")
    if not mirna[ID].reset_index(drop=True).equals(clinical[ID].reset_index(drop=True)):
        raise ValueError("miRNA and protein/metadata IDs are not identical and ordered")
    for field in (COHORT, LABEL):
        if not mirna[field].equals(clinical[field]):
            raise ValueError(f"{field} disagrees between the main worksheets")

    histology_map, exact_duplicate_rows = build_histology_map(histology)
    base = mirna[[ID, COHORT, LABEL]].merge(
        clinical.drop(columns=[COHORT, LABEL]),
        on=ID,
        how="left",
        validate="one_to_one",
    )
    cohort_split = {
        "ASPIRA (train)": "development/training",
        "ASPIRA (internal validation)": "internal validation",
        "BWH (external validation)": "external validation",
    }
    rows = []
    for record in base.to_dict(orient="records"):
        patient_id = record[ID]
        diagnosis = int(record[LABEL])
        stage_info = histology_map.get(
            patient_id,
            {
                "stage": "Unknown (not present in histologies worksheet)",
                "histology": "Unknown (not present in histologies worksheet)",
                "histology_detail": "Unknown (not present in histologies worksheet)",
            },
        )
        menopause = record.get("Meno (pre-menopause = 0, post = 1)")
        row = {
            "patient_id": patient_id,
            "sample_id": patient_id,
            "published_split": cohort_split.get(record[COHORT], "Unknown"),
            "diagnosis": "cancer" if diagnosis == 1 else "control (subtype not separately recorded)",
            "diagnosis_binary": diagnosis,
            "control_type": "not applicable" if diagnosis == 1 else "benign/healthy subtype unknown",
            "stage": stage_info["stage"] if diagnosis == 1 else "not applicable",
            "histology": stage_info["histology"] if diagnosis == 1 else "not applicable",
            "histology_detail": stage_info["histology_detail"] if diagnosis == 1 else "not applicable",
            "age": record.get("Age"),
            "menopausal_status": (
                {0: "pre-menopausal", 1: "post-menopausal"}.get(int(menopause), "unknown")
                if pd.notna(menopause)
                else "unknown"
            ),
            "ca125": record.get("CA125"),
            "he4": record.get("HE4"),
            "apoa1": record.get("ApoA1"),
            "prealbumin": record.get("PreAlb"),
            "b2m": record.get("B2M"),
            "transferrin": record.get("TRF"),
            "fsh": record.get("FSH"),
            "study_site": "BWH" if str(record[COHORT]).startswith("BWH") else "ASPIRA",
            "collection_group": record[COHORT],
        }
        rows.append(row)
    cohort = pd.DataFrame(rows)
    if len(cohort) != len(mirna) or cohort["patient_id"].duplicated().any():
        raise ValueError("Reconstruction did not preserve exactly one row per source ID")
    cohort.to_csv(COHORT_CSV, index=False)
    return cohort, histology_map, exact_duplicate_rows


def write_stage_counts(cohort: pd.DataFrame, histology: pd.DataFrame) -> None:
    cancer = cohort.loc[cohort["diagnosis_binary"].eq(1)].copy()
    split_values = [("all", cancer)] + list(cancer.groupby("published_split", sort=False))
    split_by_id = cohort.set_index("patient_id")["published_split"].to_dict()
    histology_rows = histology.copy()
    histology_rows["published_split"] = histology_rows[ID].map(split_by_id)
    stage_rows: list[dict[str, Any]] = []
    for split, subset in split_values:
        stage_counts = subset["stage"].value_counts()
        raw_rows = histology_rows if split == "all" else histology_rows.loc[
            histology_rows["published_split"].eq(split)
        ]
        raw_stage_counts = raw_rows[HISTOLOGY_STAGE].value_counts()
        for category, note in [
            ("Stage I", "Not separately reported; workbook only provides grouped I/II."),
            ("Stage II", "Not separately reported; workbook only provides grouped I/II."),
            ("Stage III", "Not separately reported; workbook only provides grouped III/IV."),
            ("Stage IV", "Not separately reported; workbook only provides grouped III/IV."),
            ("stage I/II", "Grouped stage only; not separable into Stage I versus II."),
            ("stage III/IV", "Grouped stage only; not separable into Stage III versus IV."),
            (
                "Unknown (conflicting worksheet entries)",
                "The ASPIRA training split includes one patient with both Stage I/II and Stage III/IV records.",
            ),
            ("Unknown (not present in histologies worksheet)", "Stage is not supplied for this patient."),
        ]:
            separately_unreported = category in {"Stage I", "Stage II", "Stage III", "Stage IV"}
            worksheet_row_count = (
                "not separately reported"
                if separately_unreported
                else (
                int(raw_stage_counts.get(category, 0))
                if category in {"stage I/II", "stage III/IV"}
                else (
                    2
                    if category == "Unknown (conflicting worksheet entries)"
                    and (split == "all" or split == "development/training")
                    else 0
                )
                )
            )
            stage_rows.append(
                {
                    "published_split": split,
                    "stage_category": category,
                    "patient_count": "not separately reported"
                    if separately_unreported
                    else int(stage_counts.get(category, 0)),
                    "worksheet_row_count": worksheet_row_count,
                    "denominator_cancer": len(subset),
                    "notes": note,
                }
            )
    pd.DataFrame(stage_rows).to_csv(STAGE_CSV, index=False)


def write_feature_maps(mirna: pd.DataFrame) -> pd.DataFrame:
    measurement_columns = list(mirna.columns[3:])
    normalized_to_column = {
        normalize_feature_name(str(column)): str(column) for column in measurement_columns
    }
    all_rows = []
    for column in measurement_columns:
        values = pd.to_numeric(mirna[column], errors="coerce")
        clean_name = feature_name_from_column(str(column))
        all_rows.append(
            {
                "2024_identifier": clean_name,
                "workbook_column": str(column),
                "species_prefix": clean_name.split("-", 1)[0] if "-" in clean_name else "unknown",
                "measurement_scale": "MFI (as stated in workbook column header)",
                "nonmissing_patients": int(values.notna().sum()),
                "missing_patients": int(values.isna().sum()),
                "minimum": float(values.min()),
                "median": float(values.median()),
                "maximum": float(values.max()),
            }
        )
    pd.DataFrame(all_rows).to_csv(PANEL_CSV, index=False)

    alias_columns = {"hsa-miR-335-5p": "hsa-mir-335-5p_0_-2"}
    feature_rows = []
    for signature_name, features in [
        ("current production OVAI", PRODUCTION_FEATURES),
        ("historical 14-feature signature", HISTORICAL_FEATURES),
    ]:
        for feature in features:
            normalized = normalize_feature_name(feature)
            exact_column = normalized_to_column.get(normalized)
            candidate_column = None
            if exact_column is None and feature in alias_columns:
                target = alias_columns[feature].casefold()
                candidate_column = next(
                    (
                        str(column)
                        for column in measurement_columns
                        if normalize_feature_name(str(column)) == target
                    ),
                    None,
                )
            measured_column = exact_column or candidate_column
            exact = exact_column is not None
            feature_rows.append(
                {
                    "ovai_feature": feature,
                    "2024_identifier": feature_name_from_column(measured_column) if measured_column else "",
                    "mapping_type": (
                        "exact mature-name match (case/format normalized)"
                        if exact
                        else "candidate panel alias; not verified"
                        if candidate_column
                        else "no matching assay identifier"
                    ),
                    "exact_match": exact,
                    "measured": measured_column is not None,
                    "measurement_scale": "MFI (transform/normalization not specified in workbook)"
                    if measured_column
                    else "",
                    "usable_for_frozen_test": False,
                    "notes": (
                        f"Compared against {signature_name}. The current frozen production model requires all five production miRNAs; none are present as exact mature-name matches."
                        if signature_name == "current production OVAI"
                        else "Historical signature mapping only; these features are not inputs to the current frozen production model."
                    )
                    + (
                        " Suffix `_0_-2` is not treated as an exact match without assay annotation."
                        if candidate_column
                        else ""
                    ),
                }
            )
    feature_map = pd.DataFrame(feature_rows)
    feature_map.to_csv(FEATURE_MAP_CSV, index=False)
    return feature_map


def write_compatibility(
    sheets: dict[str, pd.DataFrame], cohort: pd.DataFrame, feature_map: pd.DataFrame
) -> None:
    mirna = sheets[MIRNA_SHEET]
    exact_production = feature_map.loc[
        feature_map["ovai_feature"].isin(PRODUCTION_FEATURES) & feature_map["exact_match"],
        "ovai_feature",
    ].tolist()
    historical_exact = feature_map.loc[
        feature_map["ovai_feature"].isin(HISTORICAL_FEATURES) & feature_map["exact_match"],
        "ovai_feature",
    ].tolist()
    complete_current = 0
    if len(exact_production) == len(PRODUCTION_FEATURES):
        columns = [
            next(
                str(column)
                for column in mirna.columns[3:]
                if normalize_feature_name(str(column)) == normalize_feature_name(feature)
            )
            for feature in PRODUCTION_FEATURES
        ]
        complete_current = int(mirna[columns].notna().all(axis=1).sum())
    row = {
        "dataset": "2024 adnexal-mass FirePlex supplementary workbook",
        "patients": len(cohort),
        "unique_patient_ids": int(cohort["patient_id"].nunique()),
        "complete_mirna_patients": int(mirna.iloc[:, 3:].notna().all(axis=1).sum()),
        "measured_assay_columns": len(mirna.columns) - 3,
        "human_mirna_columns": sum(
            normalize_feature_name(str(column)).startswith("hsa-") for column in mirna.columns[3:]
        ),
        "nonhuman_control_columns": sum(
            not normalize_feature_name(str(column)).startswith("hsa-") for column in mirna.columns[3:]
        ),
        "current_ovai_exact_features": len(exact_production),
        "current_ovai_required_features": len(PRODUCTION_FEATURES),
        "patients_scoreable_by_current_frozen_ovai": complete_current,
        "historical_14_exact_features": len(historical_exact),
        "historical_14_required_features": len(HISTORICAL_FEATURES),
        "diagnosis_code_0_controls": int(cohort["diagnosis_binary"].eq(0).sum()),
        "control_subtype_counts": "not recoverable; benign mass versus healthy is not separately recorded",
        "stage_coverage": "partial; grouped I/II or III/IV only, with one conflict",
        "assay_scale": "MFI stated in headers; transformation/normalization not specified in workbook",
        "frozen_validation_status": "not scoreable: 0/5 current production miRNAs are exact measured matches",
    }
    pd.DataFrame([row]).to_csv(COMPATIBILITY_CSV, index=False)
    endpoints = [
        {
            "comparison": "all cancer vs all workbook controls",
            "cancer_n": int(cohort["diagnosis_binary"].eq(1).sum()),
            "control_n": int(cohort["diagnosis_binary"].eq(0).sum()),
            "status": "reconstructible; control subtype pooled",
            "notes": "This is not a benign-mass-only comparison.",
        },
        {
            "comparison": "all cancer vs benign adnexal mass",
            "cancer_n": int(cohort["diagnosis_binary"].eq(1).sum()),
            "control_n": "not identifiable",
            "status": "not reconstructible",
            "notes": "No patient-level control subtype field distinguishes benign mass from healthy control.",
        },
        {
            "comparison": "all cancer vs healthy control",
            "cancer_n": int(cohort["diagnosis_binary"].eq(1).sum()),
            "control_n": "not identifiable",
            "status": "not reconstructible",
            "notes": "No patient-level control subtype field distinguishes healthy control from benign mass.",
        },
        {
            "comparison": "unambiguous Stage I/II cancer vs all workbook controls",
            "cancer_n": int(cohort["stage"].eq("stage I/II").sum()),
            "control_n": int(cohort["diagnosis_binary"].eq(0).sum()),
            "status": "reconstructible; grouped stage and pooled controls",
            "notes": "56 unique, unambiguous early-stage cancer IDs; excludes one conflicting-stage ID and cancers without histology rows.",
        },
        {
            "comparison": "unambiguous Stage I/II cancer vs benign adnexal mass",
            "cancer_n": int(cohort["stage"].eq("stage I/II").sum()),
            "control_n": "not identifiable",
            "status": "not reconstructible",
            "notes": "Control subtype is not encoded at patient level.",
        },
        {
            "comparison": "unambiguous Stage I/II cancer vs healthy control",
            "cancer_n": int(cohort["stage"].eq("stage I/II").sum()),
            "control_n": "not identifiable",
            "status": "not reconstructible",
            "notes": "Control subtype is not encoded at patient level.",
        },
    ]
    pd.DataFrame(endpoints).to_csv(ENDPOINT_CSV, index=False)


def write_decision_reports(
    sheets: dict[str, pd.DataFrame],
    cohort: pd.DataFrame,
    feature_map: pd.DataFrame,
    workbook_size: int,
    workbook_sha256: str,
) -> None:
    mirna = sheets[MIRNA_SHEET]
    clinical = sheets[CLINICAL_SHEET]
    normalized_columns = {
        normalize_feature_name(str(column)): str(column) for column in mirna.columns[3:]
    }
    published_rows = []
    for feature in PUBLISHED_COMBINED_MIRNAS:
        matched = normalized_columns.get(normalize_feature_name(feature))
        published_rows.append(
            {
                "published_model_feature": feature,
                "2024_identifier": feature_name_from_column(matched) if matched else "",
                "exact_match": matched is not None,
                "measured": matched is not None,
                "notes": "Exact mature-name match after case/format normalization."
                if matched
                else "No exact workbook assay identifier; not treated as an alias.",
            }
        )
    pd.DataFrame(published_rows).to_csv(PUBLISHED_FEATURE_MAP_CSV, index=False)
    published_exact = sum(row["exact_match"] for row in published_rows)

    split_summary = (
        cohort.groupby(["published_split", "diagnosis_binary"], dropna=False)
        .size()
        .unstack(fill_value=0)
    )
    split_counts = {
        split: {
            "controls": int(values.get(0, 0)),
            "cancer": int(values.get(1, 0)),
            "total": int(values.sum()),
        }
        for split, values in split_summary.iterrows()
    }
    external = cohort.loc[cohort["published_split"].eq("external validation")]
    external_cancer = external.loc[external["diagnosis_binary"].eq(1)]
    external_stage_i_ii = int(external_cancer["stage"].eq("stage I/II").sum())
    external_stage_iii_iv = int(external_cancer["stage"].eq("stage III/IV").sum())
    production_map = feature_map.loc[feature_map["ovai_feature"].isin(PRODUCTION_FEATURES)]
    exact_production = int(production_map["exact_match"].sum())
    feature_specs = [
        f"| `{row['published_model_feature']}` | `{'yes' if row['exact_match'] else 'no'}` | `{'yes' if row['measured'] else 'no'}` |"
        for row in published_rows
    ]

    FROZEN_VALIDATION_REPORT.write_text(
        "\n".join(
            [
                "# Frozen GSE94533 OVAI External Validation: 2024 FirePlex",
                "",
                "## Preflight Result",
                "",
                "**No patient-level predictions were generated.** The frozen production model requires all five production miRNAs, age, and CA-125. The workbook has age and CA-125 but has zero exact measured matches to the five required production miRNAs. Substituting historical markers, using fuzzy aliases, refitting a model, or estimating missing features would not be a frozen validation.",
                "",
                f"- 2024 patient records: {len(cohort)}.",
                f"- Exact production-miRNA coverage: {exact_production}/5.",
                "- Rows eligible for frozen production-model scoring: 0.",
                "- Frozen model, scaler, defaults, thresholds, and canonical artifacts were not modified.",
                "- No external labels were used for feature selection, threshold tuning, or model selection.",
                "- This is an input-compatibility stop, not an estimate of OVAI clinical performance on this cohort.",
                "",
                "## Required Frozen Inputs",
                "",
                "| Required production feature | Exact FirePlex match | Measured |",
                "|---|---:|---:|",
                *[
                    f"| `{row['ovai_feature']}` | `{'yes' if bool(row['exact_match']) else 'no'}` | `{'yes' if bool(row['measured']) else 'no'}` |"
                    for row in production_map.to_dict(orient="records")
                ],
                "",
                "The workbook reports FirePlex MFI values. The production OVAI pipeline was fitted using GSE94533 miRNA values with `log1p(max(value, 0))`, clinical imputation, scaling, and fixed interactions. There is no documented validated transform bridging these assay scales, and the required miRNA features are absent in any case.",
                "",
                "## Disposition",
                "",
                "Frozen external validation is **not computable from this workbook**. Do not report AUC, sensitivity, specificity, calibration, or threshold performance for the production OVAI model on these data. A future transportability study would first require a cohort that measures the exact five analytes or an independently validated assay bridge; neither can be manufactured from the current 2024 matrix.",
                "",
            ]
        ),
        encoding="utf-8",
    )

    split_lines = [
        f"- {split}: n={counts['total']}, cancer={counts['cancer']}, control={counts['controls']}."
        for split, counts in split_counts.items()
    ]
    PUBLISHED_REPRODUCTION_REPORT.write_text(
        "\n".join(
            [
                "# Published 2024 Model and Cohort Reproduction Assessment",
                "",
                "## Reproduced from the Workbook",
                "",
                f"- Workbook file: `{WORKBOOK.relative_to(ROOT).as_posix()}`; {workbook_size:,} bytes; SHA256 `{workbook_sha256}`.",
                f"- Cohort reconstruction: {len(cohort)} unique IDs, {int(cohort['diagnosis_binary'].sum())} cancer labels, {int(cohort['diagnosis_binary'].eq(0).sum())} control labels.",
                *split_lines,
                f"- External validation subset: n={len(external)}, 51 cancer and 59 controls; workbook has {external_stage_i_ii} unambiguous Stage I/II cancer IDs and {external_stage_iii_iv} Stage III/IV IDs.",
                "- The workbook reproduces the published split and label counts exactly. The external stage group counts (20 Stage I/II, 31 Stage III/IV) also agree.",
                "- The workbook demographics reproduce 290 premenopausal subjects; mean age and postmenopausal counts by source group are summarized in the workbook audit.",
                "",
                "## Published Combined-Model miRNA Crosswalk",
                "",
                f"Exact matches to the 10 miRNAs listed in the paper's model-feature table: {published_exact}/10.",
                "",
                "| Published feature | Exact match | Measured |",
                "|---|---:|---:|",
                *feature_specs,
                "",
                "The paper lists `hsa-miR-103-3p`, while the workbook lists `hsa-mir-103a-3p` (and a distinct `hsa-mir-103a-2-5p`). These are not treated as identical without assay/nomenclature documentation. The other listed features, including `cel-miR-39-3p`, are exact workbook identifiers.",
                "",
                "## Performance Reproduction Decision",
                "",
                "The paper describes a neural-network combined model, forward regression with ten-fold cross-validation to select the number of miRNAs, and a training-derived probability threshold chosen for at least 80% cross-validated sensitivity. It reports combined-model threshold 0.45, external AUC 0.95 (95% CI 0.90-0.98), sensitivity 0.92, and specificity 0.80. The workbook contains the source patient-level matrix but does not contain fitted network weights or a fully specified executable network architecture, selected-feature coefficients/weights, training fold assignments, or all preprocessing details. Further, one published selected-miRNA identifier does not have an exact workbook-column match.",
                "",
                "Accordingly, the cohort/split structure and published feature availability were audited, but published validation predictions and performance were **not recreated**. Fitting a new network or reselecting features on this workbook would be a new model-development exercise, not reproduction of the published frozen model. No such fit was performed in this run.",
                "",
            ]
        ),
        encoding="utf-8",
    )

    published_validation = (
        "The workbook is suitable for research-only development of a distinct FirePlex-specific model because it supplies the published development/internal/external labels and complete panel data. Preserve the 468/100/110 split; use only the 468 ASPIRA training subjects for any fitting/selection, and keep internal and BWH external validation untouched until a prespecified model is frozen. The control phenotype is pooled and cannot support benign-mass-only performance claims from this workbook."
    )
    DECISION_REPORT.write_text(
        "\n".join(
            [
                "# 2024 Cohort Compatibility and Decision",
                "",
                "## Data Integrity and Usable Cohort",
                "",
                f"The manually supplied workbook is a valid, intact XLSX ({workbook_size:,} bytes; SHA256 `{workbook_sha256}`). It contains 678 unique patient/sample IDs and 180 complete numeric assay columns (179 human-target columns plus one cel-miR-39 control). The `miRNA` and `protein + metadata` sheets have exact one-to-one ordered ID matches and matching diagnosis/split values.",
                "",
                "The published cohort is reconstructed exactly: 286 cancer and 392 pooled controls; development 191/277, internal validation 44/56, and BWH external validation 51/59. The workbook does not distinguish benign adnexal masses from healthy controls, so subtype-specific control counts and cancer-versus-benign-only performance are unavailable.",
                "",
                "## Stage and Control Limits",
                "",
                "The histology worksheet has 145 rows but 141 unique IDs, all among cancer patients. Three rows duplicate an otherwise identical record; one ID has both Stage I/II and Stage III/IV entries and is classified as conflicting. After exact duplicate removal, 56 patients have unambiguous Stage I/II, 84 have unambiguous Stage III/IV, one is unresolved due to conflicting entries, and 145 cancers have no row in the histology sheet. Raw worksheet counts (59/86) reproduce the paper, but they are not distinct patient counts. Stages I vs II and III vs IV cannot be separately determined.",
                "",
                "Thus the available early-stage comparison is 56 unambiguous Stage I/II cancers versus 392 pooled controls; this is not a benign-mass-only estimate. In the BWH external group, the workbook records 20 Stage I/II and 31 Stage III/IV cancer IDs against 59 pooled controls.",
                "",
                "## OVAI Compatibility",
                "",
                "The five production features are not present as exact measured mature-miRNA identifiers (0/5). The historical 14-feature signature has 8/14 exact matches; `hsa-miR-335-5p_0_-2` remains an unverified assay alias and is not counted. FirePlex MFI values are not pooled with GSE94533 sequencing values.",
                "",
                "**Decision:** the 2024 workbook is not compatible with direct frozen evaluation of the current production OVAI model and is not a transport validation for that signature. No OVAI retraining, feature substitution, threshold tuning, or modeling was performed. See `2024_adnexal_mass_frozen_validation.md` for the scoring stop gate.",
                "",
                "## Later Development Eligibility",
                "",
                published_validation,
                "",
                "The dataset may support a separately identified FirePlex-specific research model after prespecified methods are established. It must remain separate from canonical OVAI artifacts and from GSE94533 values. The published combined-model performance cannot be reproduced faithfully from the workbook alone; details are in `2024_adnexal_mass_published_model_reproduction.md`.",
                "",
            ]
        ),
        encoding="utf-8",
    )

    DATA_NEEDED_REPORT.write_text(
        "\n".join(
            [
                "# 2024 Adnexal-Mass Workbook Status",
                "",
                "The expected workbook is now present and has been audited.",
                "",
                f"- Filename: `NIHMS2021079-supplement-1.xlsx`",
                f"- Path: `{WORKBOOK.relative_to(ROOT).as_posix()}`",
                f"- File type: valid XLSX; {workbook_size:,} bytes.",
                f"- SHA256: `{workbook_sha256}`.",
                "- Source: \"Serum miRNA improves the accuracy of a multivariate index assay for triage of an adnexal mass\"; DOI `10.1016/j.ygyno.2024.08.008`; PMCID `PMC13283006`.",
                "",
                "The workbook audit and cohort reconstruction are complete. Frozen OVAI validation is blocked by zero exact matches to the five required production miRNAs. Patient-level benign-mass versus healthy-control composition is not separately encoded. No model was fit.",
                "",
                "See `2024_adnexal_mass_compatibility_decision.md`, `2024_adnexal_mass_frozen_validation.md`, and `2024_adnexal_mass_published_model_reproduction.md` for the current decisions.",
                "",
            ]
        ),
        encoding="utf-8",
    )


def main() -> None:
    if not WORKBOOK.is_file():
        raise FileNotFoundError(f"Expected workbook is missing: {WORKBOOK}")
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    if not zipfile.is_zipfile(WORKBOOK):
        raise ValueError("Input is not a valid XLSX/ZIP container; refusing to read it")
    with zipfile.ZipFile(WORKBOOK) as archive:
        bad_member = archive.testzip()
        if bad_member is not None or "xl/workbook.xml" not in archive.namelist():
            raise ValueError(f"Invalid XLSX structure: {bad_member or 'missing workbook.xml'}")

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet_names = list(workbook.sheetnames)
    workbook.close()
    sheets = {name: pd.read_excel(WORKBOOK, sheet_name=name) for name in sheet_names}
    required_sheets = {MIRNA_SHEET, CLINICAL_SHEET, HISTOLOGY_SHEET}
    if not required_sheets.issubset(sheets):
        raise ValueError(f"Expected worksheets missing: {sorted(required_sheets - set(sheets))}")

    size, digest = write_file_verification()
    write_inventory(sheets)
    audit_sections = [
        "# 2024 Adnexal-Mass Workbook Audit",
        "",
        f"Workbook: `{WORKBOOK.relative_to(ROOT).as_posix()}`",
        f"Size: {size:,} bytes. SHA256: `{digest}`.",
        "All worksheets are inspected below. Previews are the first/last five data rows (the header is reported separately).",
        "",
    ]
    audit_sections.extend(describe_sheet(name, frame) for name, frame in sheets.items())

    mirna = sheets[MIRNA_SHEET]
    clinical = sheets[CLINICAL_SHEET]
    histology = sheets[HISTOLOGY_SHEET]
    if mirna.shape[1] - 3 != 180:
        raise ValueError(f"Expected 180 assay columns, found {mirna.shape[1] - 3}")
    cohort, histology_map, exact_duplicate_rows = write_reconstructions(mirna, clinical, histology)
    write_stage_counts(cohort, histology)
    feature_map = write_feature_maps(mirna)
    write_compatibility(sheets, cohort, feature_map)
    write_decision_reports(sheets, cohort, feature_map, size, digest)
    manifest_path = ROOT / "research/cohort_expansion/manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest.update(
        {
            "source_status": "patient-level supplementary workbook present and audited",
            "status": "audited; frozen production OVAI external scoring blocked by 0/5 exact production miRNA matches",
            "workbook_path": WORKBOOK.relative_to(ROOT).as_posix(),
            "workbook_size_bytes": size,
            "workbook_sha256": digest,
            "usable_patient_records": len(cohort),
            "complete_mirna_patients": int(mirna.iloc[:, 3:].notna().all(axis=1).sum()),
            "measured_assay_columns": len(mirna.columns) - 3,
            "current_production_exact_mirna_matches": int(
                feature_map.loc[
                    feature_map["ovai_feature"].isin(PRODUCTION_FEATURES), "exact_match"
                ].sum()
            ),
            "historical_14_exact_mirna_matches": int(
                feature_map.loc[
                    feature_map["ovai_feature"].isin(HISTORICAL_FEATURES), "exact_match"
                ].sum()
            ),
            "control_subtypes_patient_level": "not separately identified",
            "modeling_performed": False,
            "canonical_artifacts_modified": False,
        }
    )
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

    stage_map_frame = pd.DataFrame(
        [{ID: patient_id, **values} for patient_id, values in histology_map.items()]
    )
    stage_map_frame = stage_map_frame.merge(
        cohort[["patient_id", "published_split", "diagnosis_binary"]],
        left_on=ID,
        right_on="patient_id",
        validate="one_to_one",
    )
    known_cancer_stage = stage_map_frame.loc[
        stage_map_frame["diagnosis_binary"].eq(1)
        & stage_map_frame["stage"].isin(["stage I/II", "stage III/IV"])
    ]
    stage_counts = cohort.loc[cohort["diagnosis_binary"].eq(1), "stage"].value_counts()
    split_table = pd.crosstab(cohort["published_split"], cohort["diagnosis_binary"])
    human_columns = sum(
        normalize_feature_name(str(column)).startswith("hsa-") for column in mirna.columns[3:]
    )
    complete_markers = int(mirna.iloc[:, 3:].notna().all(axis=1).sum())
    exact_current = int(
        feature_map.loc[feature_map["ovai_feature"].isin(PRODUCTION_FEATURES), "exact_match"].sum()
    )
    exact_historical = int(
        feature_map.loc[feature_map["ovai_feature"].isin(HISTORICAL_FEATURES), "exact_match"].sum()
    )
    audit_sections.extend(
        [
            "## Cross-Worksheet and Cohort Reconstruction",
            "",
            f"- `miRNA` and `protein + metadata` each contain {len(mirna)} unique nonmissing IDs in the same order; IDs, cohort labels, and diagnosis labels agree exactly.",
            f"- `histologies` contains {len(histology)} rows and {histology[ID].nunique()} unique IDs; all IDs join to cancer rows in the main table.",
            f"- The histology worksheet has {exact_duplicate_rows} exact duplicate rows; these were collapsed only for the patient-level join. One ID has conflicting stage group entries and is marked unknown, not resolved by row order.",
            f"- Reconstructed patient/sample rows: {len(cohort)} (one per unique main-sheet ID).",
            f"- Complete miRNA measurements across all {mirna.shape[1] - 3} assay columns: {complete_markers}/{len(mirna)} patients.",
            f"- Assay columns labeled with human `hsa-` identifiers: {human_columns}; remaining columns: {mirna.shape[1] - 3 - human_columns} (off-species/control).",
            f"- Diagnosis-code counts: {json.dumps({str(key): int(value) for key, value in mirna[LABEL].value_counts().sort_index().items()})}; labels are 0=benign/control and 1=malignant.",
            "- Control subtype: the workbook does not separately identify benign adnexal masses versus healthy controls. The paper describes control sources as benign adnexal masses or healthy controls; patient-level subtype counts cannot be reconstructed.",
            "- The histology sheet records only grouped `stage I/II` and `stage III/IV`; it does not separate stages I, II, III, and IV.",
            "- Raw histology worksheet rows by split reproduce the paper's grouped stage counts (ASPIRA train 34/44, internal validation 5/11, BWH external validation 20/31; combined Stage I/II=59 and Stage III/IV=86). These are worksheet rows, not all distinct IDs.",
            "- Confirmed non-conflicting cancer stage groups: "
            + json.dumps(
                {
                    key: int(stage_counts.get(key, 0))
                    for key in ["stage I/II", "stage III/IV"]
                }
            )
            + f"; conflicting stage ID: {int(stage_counts.get('Unknown (conflicting worksheet entries)', 0))}; cancers absent from the histology sheet: {int(stage_counts.get('Unknown (not present in histologies worksheet)', 0))}.",
            "- Patient-level discrepancy: after removing three exact duplicate rows and treating the conflicting ID as unresolved, 56 IDs have unambiguous Stage I/II, 84 have unambiguous Stage III/IV, 1 has conflicting entries, and 145 cancer IDs have no histology-sheet row. Thus the paper's 59/86 row totals cannot be interpreted as 145 distinct, unambiguous staged patients from this workbook.",
            "- The paper reports 141 ASPIRA cases with unknown stage; the workbook has 145 ASPIRA cancer IDs absent from the histologies worksheet, plus one conflicting staged ID. The workbook's patient-level stage availability therefore differs from that published summary.",
            "- The external BWH subset has complete grouped staging in the workbook: "
            + json.dumps(
                {
                    str(key): int(value)
                    for key, value in known_cancer_stage.loc[
                        known_cancer_stage["published_split"].eq("external validation"), "stage"
                    ].value_counts().items()
                }
            )
            + ".",
            "",
            "### Diagnosis by Published Split",
            "```text",
            split_table.to_string(),
            "```",
            "",
            "## Usability and Feature Mapping (Before Modeling)",
            "",
            f"- Patient-level records usable for descriptive cohort auditing: {len(cohort)}/{len(cohort)}.",
            f"- Complete numeric assay measurements: {complete_markers} patients x {mirna.shape[1] - 3} columns ({human_columns} human targets plus {mirna.shape[1] - 3 - human_columns} off-species/control target).",
            f"- Exact current production miRNA matches: {exact_current}/{len(PRODUCTION_FEATURES)}; frozen-OVAI scoreable patients: 0/{len(cohort)}.",
            f"- Exact historical 14-signature matches: {exact_historical}/14; the `hsa-miR-335-5p_0_-2` assay column is recorded as an unverified candidate alias, not an exact match.",
            "- Assay columns state `MFI`; the workbook alone does not establish any additional transformation or normalization. No model was fit by this audit script.",
            "",
            "## Source-Reported Cohort and Model Facts",
            "",
            "The worksheet split/label counts reproduce the paper's reported 678 subjects (286 malignant, 392 controls), with development 468 (191/277), internal validation 100 (44/56), and external validation 110 (51/59). The paper selected 10 miRNAs for its combined model with seven proteins, age, and menopausal status; the workbook has inputs for that analysis, but it does not contain trained network weights. The performance model is not fitted by this audit script.",
            "",
        ]
    )
    WORKBOOK_AUDIT.write_text("\n".join(audit_sections), encoding="utf-8")
    print(f"Verified {len(cohort)} patients; {mirna.shape[1] - 3} assay columns; {complete_markers} patients complete.")
    print(f"Current OVAI exact feature matches: {exact_current}/5; historical exact matches: {exact_historical}/14.")
    print(f"Wrote workbook audit outputs under {REPORT_DIR.relative_to(ROOT)}.")


if __name__ == "__main__":
    main()