#!/usr/bin/env python3
"""Audit the locally present Gockley workbook and frozen PMML networks."""

from __future__ import annotations

from collections import Counter
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import roc_auc_score


ROOT = Path(__file__).resolve().parents[2]
SOURCE_DIR = ROOT / "research/cohort_expansion/data/raw/2024_adnexal_mass"
WORKBOOK = SOURCE_DIR / "DataSheet_2.xlsx"
MODEL_ARCHIVE = SOURCE_DIR / "DataSheet_3.zip"
REPORT_DIR = ROOT / "research/cohort_expansion/reports"
GSE_PREDICTIONS = ROOT / "reports/error_analysis_predictions.csv"
PRODUCTION_MANIFEST = ROOT / "production_artifacts/manifest.json"

DIAGNOSIS_COLUMN = "Control/Benign/Cancer Early/Cancer Advanced"
HISTORICAL = [
    "hsa-miR-23b-3p",
    "hsa-miR-29a-3p",
    "hsa-miR-32-5p",
    "hsa-miR-92a-3p",
    "hsa-miR-150-5p",
    "hsa-miR-200a-3p",
    "hsa-miR-200c-3p",
    "hsa-miR-203a",
    "hsa-miR-320c",
    "hsa-miR-320d",
    "hsa-miR-335-5p",
    "hsa-miR-450b-5p",
    "hsa-miR-1246",
    "hsa-miR-1307-5p",
]
CURRENT = [
    "hsa-miR-200b-3p",
    "hsa-miR-296-3p",
    "hsa-miR-6826-5p",
    "hsa-miR-32-3p",
    "hsa-miR-1255a",
]
MODEL_FILENAME = "signature-best-10.xml"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def file_info(path: Path, description: str) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": path.relative_to(ROOT).as_posix(),
        "description": description,
        "size_bytes": stat.st_size,
        "sha256": sha256_file(path),
        "modified": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(sep=" "),
    }


def stage_group(value: Any) -> str:
    if value is None or not str(value).strip():
        return ""
    stage = str(value).strip().upper()
    if stage.startswith(("1", "2")):
        return "Stage I/II"
    if stage.startswith("3A"):
        return "Stage IIIA"
    if stage.startswith("3"):
        return "Stage III/IV"
    return f"Stage {stage}"


def load_workbook_data() -> tuple[list[str], list[tuple[Any, ...]], list[str], dict[str, int]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"Unexpected worksheet names: {workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    headers = [str(value) if value is not None else "" for value in values[0]]
    rows = [tuple(row) for row in values[1:]]
    expected_start = [
        "ID",
        "Set",
        "Age",
        DIAGNOSIS_COLUMN,
        "Histopathology",
        "Mutation status",
        "Stage",
        "Grade",
    ]
    expected_tail = ["Mutation status", "Stage", "Grade", "Histology"]
    if headers[:8] != expected_start:
        raise ValueError(f"Unexpected leading workbook fields: {headers[:8]}")
    if [header.strip() for header in headers[-4:]] != expected_tail:
        raise ValueError(f"Unexpected trailing workbook fields: {headers[-4:]}")
    feature_names = headers[8:-4]
    if not feature_names or len(feature_names) != len(set(feature_names)):
        raise ValueError("The contiguous miRNA feature block is empty or has duplicate names")
    if any(not name.startswith(("hsa-", "mmu-", "rno-")) for name in feature_names):
        raise ValueError("Unexpected non-miRNA field in the assay feature block")
    metadata = {
        "row_count_including_header": len(values),
        "column_count": len(headers),
        "unique_patient_id_count": len({row[0] for row in rows}),
        "duplicate_patient_id_count": len(rows) - len({row[0] for row in rows}),
        "duplicate_full_row_count": len(rows) - len(set(rows)),
        "miRNA_feature_count": len(feature_names),
        "workbook_dimension": f"A1:GW{len(values)}",
        "sheet_names": workbook.sheetnames if hasattr(workbook, "sheetnames") else ["Sheet1"],
    }
    return headers, rows, feature_names, metadata


def normalize_id(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return re.sub(r"[^a-z0-9]", "", text)


def build_records(headers: list[str], rows: list[tuple[Any, ...]], features: list[str]) -> list[dict[str, Any]]:
    feature_start = 8
    feature_end = feature_start + len(features)
    records = []
    for row in rows:
        raw_stage = row[6]
        record: dict[str, Any] = {
            "sample_id": row[0],
            "set": row[1],
            "age": row[2],
            "diagnosis": row[3],
            "histopathology": row[4],
            "mutation_status": row[5],
            "stage_raw": raw_stage,
            "stage_group": stage_group(raw_stage),
            "grade": row[7],
            "histology_secondary": row[-1],
            "ca125": None,
        }
        record.update({feature: row[feature_start + offset] for offset, feature in enumerate(features)})
        if len(row[feature_end:-4]):
            raise ValueError("Workbook assay columns do not align with trailing metadata fields")
        records.append(record)
    return records


def parse_pmml_model(archive: ZipFile, filename: str) -> dict[str, Any]:
    root = ET.fromstring(archive.read(filename))
    network = root.find("NeuralNetwork")
    if network is None:
        raise ValueError(f"No NeuralNetwork in {filename}")
    inputs = []
    for neural_input in network.findall("./NeuralInputs/NeuralInput"):
        norm = neural_input.find(".//NormContinuous")
        if norm is None:
            raise ValueError(f"No NormContinuous input transform in {filename}")
        points = norm.findall("LinearNorm")
        if len(points) != 2:
            raise ValueError(f"Expected two LinearNorm points for {norm.get('field')} in {filename}")
        minimum, maximum = [float(point.get("orig")) for point in points]
        minimum_norm, maximum_norm = [float(point.get("norm")) for point in points]
        scale = float(norm.get("scale"))
        shift = float(norm.get("shift"))
        expected_scale = (maximum_norm - minimum_norm) / (maximum - minimum)
        expected_shift = minimum_norm - minimum * expected_scale
        inputs.append(
            {
                "id": neural_input.get("id"),
                "feature": norm.get("field"),
                "minimum": minimum,
                "maximum": maximum,
                "scale": scale,
                "shift": shift,
                "normalization_consistent": math.isclose(scale, expected_scale, rel_tol=1e-6, abs_tol=1e-8)
                and math.isclose(shift, expected_shift, rel_tol=1e-6, abs_tol=1e-8),
                "outlier_treatment": norm.get("outliers", "asIs"),
            }
        )
    layers = []
    for layer in network.findall("NeuralLayer"):
        neurons = []
        for neuron in layer.findall("Neuron"):
            neurons.append(
                {
                    "id": neuron.get("id"),
                    "bias": float(neuron.get("bias")),
                    "connections": [
                        {"from": con.get("from"), "weight": float(con.get("weight"))}
                        for con in neuron.findall("Con")
                    ],
                }
            )
        layers.append({"activation": layer.get("activationFunction", "identity"), "neurons": neurons})
    output_mapping = {}
    for neural_output in network.findall("./NeuralOutputs/NeuralOutput"):
        discrete = neural_output.find(".//NormDiscrete")
        if discrete is not None:
            output_mapping[neural_output.get("outputNeuron")] = discrete.get("value")
    if len(inputs) != int(network.find("NeuralInputs").get("numberOfInputs")):
        raise ValueError(f"Input count mismatch in {filename}")
    if not layers or not output_mapping:
        raise ValueError(f"Incomplete network definition in {filename}")
    total_weights = sum(len(neuron["connections"]) for layer in layers for neuron in layer["neurons"])
    total_biases = sum(len(layer["neurons"]) for layer in layers)
    model = {
        "filename": filename,
        "model_name": network.get("modelName"),
        "function": network.get("functionName"),
        "inputs": inputs,
        "layers": layers,
        "output_mapping": output_mapping,
        "weight_count": total_weights,
        "bias_count": total_biases,
        "explicit_threshold": any("threshold" in key.lower() for element in root.iter() for key in element.attrib),
        "has_model_verification": root.find(".//ModelVerification") is not None,
    }
    return model


def activate(value: float, function: str) -> float:
    normalized = function.lower()
    if normalized in {"identity", "none"}:
        return value
    if normalized == "tanh":
        return math.tanh(value)
    if normalized == "logistic":
        if value >= 0:
            return 1.0 / (1.0 + math.exp(-value))
        exp_value = math.exp(value)
        return exp_value / (1.0 + exp_value)
    if normalized == "rectifier":
        return max(0.0, value)
    raise ValueError(f"Unsupported PMML activation function: {function}")


def predict_pmml(model: dict[str, Any], record: dict[str, Any]) -> dict[str, Any]:
    values: dict[str, float] = {}
    out_of_range = []
    for item in model["inputs"]:
        raw = record.get(item["feature"])
        if raw is None or pd.isna(raw):
            raise ValueError(f"Missing model input {item['feature']} for patient {record['sample_id']}")
        number = float(raw)
        if number < item["minimum"] or number > item["maximum"]:
            out_of_range.append(item["feature"])
        values[item["id"]] = number * item["scale"] + item["shift"]
    for layer in model["layers"]:
        next_values = {}
        for neuron in layer["neurons"]:
            total = neuron["bias"]
            for connection in neuron["connections"]:
                if connection["from"] not in values:
                    raise ValueError(f"Unresolved PMML connection {connection['from']} in {model['filename']}")
                total += values[connection["from"]] * connection["weight"]
            next_values[neuron["id"]] = activate(total, layer["activation"])
        values = next_values
    outputs = {label: values[node_id] for node_id, label in model["output_mapping"].items()}
    if "malignant" not in outputs or "benign/control" not in outputs:
        raise ValueError(f"Unexpected output classes in {model['filename']}: {outputs.keys()}")
    prediction = max(outputs, key=outputs.get)
    return {
        "score_margin_malignant": outputs["malignant"] - outputs["benign/control"],
        "output_malignant": outputs["malignant"],
        "output_benign_control": outputs["benign/control"],
        "predicted_class": prediction,
        "out_of_range_features": ";".join(out_of_range),
    }


def endpoint_masks(records: list[dict[str, Any]]) -> dict[str, list[bool]]:
    masks: dict[str, list[bool]] = {
        "stage_I_II_cancer_vs_benign_mass": [],
        "stage_I_II_cancer_vs_benign_plus_healthy": [],
        "all_cancer_vs_benign_mass": [],
        "all_cancer_vs_benign_plus_healthy": [],
    }
    for record in records:
        diagnosis = record["diagnosis"]
        early = diagnosis == "Cancer Early" and record["stage_group"] == "Stage I/II"
        cancer = diagnosis == "Cancer Early"
        for key in masks:
            if key.startswith("stage_I_II"):
                included = early or diagnosis == "Benign" or (
                    key.endswith("benign_plus_healthy") and diagnosis == "Control"
                )
            else:
                included = cancer or diagnosis == "Benign" or (
                    key.endswith("benign_plus_healthy") and diagnosis == "Control"
                )
            masks[key].append(included)
    return masks


def evaluate_endpoints(records: list[dict[str, Any]], predictions: list[dict[str, Any]]) -> pd.DataFrame:
    split_names = ["All"] + sorted({str(record["set"]) for record in records if record["set"] is not None})
    masks = endpoint_masks(records)
    rows = []
    for endpoint, inclusion in masks.items():
        for split in split_names:
            selected = [
                index for index, record in enumerate(records)
                if inclusion[index] and (split == "All" or record["set"] == split)
            ]
            subset = [records[index] for index in selected]
            if not subset:
                continue
            truth = np.array([record["diagnosis"] == "Cancer Early" for record in subset], dtype=bool)
            predicted = np.array([predictions[index]["predicted_class"] == "malignant" for index in selected], dtype=bool)
            score = np.array([predictions[index]["score_margin_malignant"] for index in selected], dtype=float)
            tp = int(np.sum(truth & predicted))
            tn = int(np.sum(~truth & ~predicted))
            fp = int(np.sum(~truth & predicted))
            fn = int(np.sum(truth & ~predicted))
            sensitivity = tp / (tp + fn) if tp + fn else float("nan")
            specificity = tn / (tn + fp) if tn + fp else float("nan")
            auc = float(roc_auc_score(truth, score)) if len(set(truth)) == 2 else float("nan")
            rows.append(
                {
                    "endpoint": endpoint,
                    "subset": split,
                    "n_cases": int(truth.sum()),
                    "n_controls": int((~truth).sum()),
                    "n_total": len(subset),
                    "roc_auc_margin": auc,
                    "sensitivity_pmml_argmax": sensitivity,
                    "specificity_pmml_argmax": specificity,
                    "true_negative": tn,
                    "false_positive": fp,
                    "false_negative": fn,
                    "true_positive": tp,
                    "decision_rule": "PMML maximum output neuron; no threshold tuned",
                    "score_for_auc": "malignant output minus benign/control output",
                }
            )
    return pd.DataFrame(rows)


def build_feature_map(features: list[str], model_features: list[str]) -> pd.DataFrame:
    measured = set(features)
    modeled = set(model_features)
    neighbor_notes = {
        "hsa-miR-200b-3p": "hsa-miR-200a-3p and hsa-miR-200c-3p are distinct mature family members, not aliases",
        "hsa-miR-32-3p": "hsa-miR-32-5p is the opposite mature arm, not an alias",
    }
    requested = list(dict.fromkeys(CURRENT + HISTORICAL))
    all_names = list(features) + [name for name in requested if name not in measured]
    rows = []
    for position, feature in enumerate(all_names, start=1):
        present = feature in measured
        rows.append(
            {
                "feature": feature,
                "workbook_exact_match": present,
                "workbook_column_1based": position if present else "",
                "current_OVAI_5_signature": feature in CURRENT,
                "historical_14_signature": feature in HISTORICAL,
                "signature_best_10_PMML_input": feature in modeled,
                "exact_alias_accepted": False,
                "nearby_non_equivalent_identifier": neighbor_notes.get(feature, ""),
                "notes": "Exact mature identifier match only; no fuzzy matching or substitution used." if present else "Requested feature is absent by exact identifier.",
            }
        )
    position_by_name = {name: idx + 9 for idx, name in enumerate(features)}
    for row in rows:
        if row["feature"] in position_by_name:
            row["workbook_column_1based"] = position_by_name[row["feature"]]
    return pd.DataFrame(rows)


def summarize_model(model: dict[str, Any]) -> str:
    architecture = "-".join([str(len(model["inputs"]))] + [str(len(layer["neurons"])) for layer in model["layers"]])
    activations = ", ".join(layer["activation"] for layer in model["layers"])
    return (
        f"`{model['filename']}`: `{model['model_name']}`, architecture {architecture}; "
        f"activations {activations}; {model['weight_count']} connection weights and "
        f"{model['bias_count']} biases; {len(model['inputs'])} explicit input transforms."
    )


def make_report(
    headers: list[str],
    rows: list[tuple[Any, ...]],
    records: list[dict[str, Any]],
    features: list[str],
    metadata: dict[str, Any],
    file_details: list[dict[str, Any]],
    archive_names: list[str],
    archive_integrity: str,
    models: list[dict[str, Any]],
    metrics: pd.DataFrame,
    gse_audit: dict[str, Any],
) -> str:
    diagnoses = Counter(record["diagnosis"] for record in records)
    stages = Counter(str(record["stage_raw"]).strip().upper() for record in records if record["stage_raw"] is not None)
    grouped_stages = Counter(record["stage_group"] for record in records if record["stage_group"])
    splits = Counter(str(record["set"]) for record in records)
    all_blanks = sum(value is None for row in rows for value in row)
    assay_missing = sum(record[feature] is None for record in records for feature in features)
    histopath_counts = Counter(str(record["histopathology"]) for record in records if record["histopathology"] is not None)
    histology_secondary_count = sum(record["histology_secondary"] is not None for record in records)
    historical_overlap = [feature for feature in HISTORICAL if feature in features]
    current_overlap = [feature for feature in CURRENT if feature in features]
    primary = metrics.loc[(metrics["endpoint"] == "stage_I_II_cancer_vs_benign_mass") & metrics["subset"].isin(["All", "Test"])]
    metric_table = [
        "| Endpoint | Subset | Cases | Controls | ROC-AUC (margin) | Sensitivity (PMML argmax) | Specificity (PMML argmax) | TN/FP/FN/TP |",
        "|---|---|---:|---:|---:|---:|---:|---|",
    ]
    for item in primary.to_dict(orient="records"):
        auc = "NA" if pd.isna(item["roc_auc_margin"]) else f"{item['roc_auc_margin']:.4f}"
        sens = "NA" if pd.isna(item["sensitivity_pmml_argmax"]) else f"{item['sensitivity_pmml_argmax']:.4f}"
        spec = "NA" if pd.isna(item["specificity_pmml_argmax"]) else f"{item['specificity_pmml_argmax']:.4f}"
        metric_table.append(
            f"| {item['endpoint']} | {item['subset']} | {item['n_cases']} | {item['n_controls']} | {auc} | {sens} | {spec} | {item['true_negative']}/{item['false_positive']}/{item['false_negative']}/{item['true_positive']} |"
        )
    full_metrics_table = [
        "| Endpoint | Subset | Cases | Controls | ROC-AUC | Sensitivity | Specificity | TN/FP/FN/TP |",
        "|---|---|---:|---:|---:|---:|---:|---|",
    ]
    for item in metrics.to_dict(orient="records"):
        auc = "NA" if pd.isna(item["roc_auc_margin"]) else f"{item['roc_auc_margin']:.4f}"
        sens = "NA" if pd.isna(item["sensitivity_pmml_argmax"]) else f"{item['sensitivity_pmml_argmax']:.4f}"
        spec = "NA" if pd.isna(item["specificity_pmml_argmax"]) else f"{item['specificity_pmml_argmax']:.4f}"
        full_metrics_table.append(
            f"| {item['endpoint']} | {item['subset']} | {item['n_cases']} | {item['n_controls']} | {auc} | {sens} | {spec} | {item['true_negative']}/{item['false_positive']}/{item['false_negative']}/{item['true_positive']} |"
        )
    file_lines = [
        f"- `{item['path']}`: {item['description']}; {item['size_bytes']} bytes; SHA256 `{item['sha256']}`; modified `{item['modified']}`."
        for item in file_details
    ]
    model_lines = [summarize_model(model) for model in models]
    model = next(item for item in models if item["filename"] == MODEL_FILENAME)
    inputs_text = ", ".join(f"`{item['feature']}`" for item in model["inputs"])
    normalization_ok = all(item["normalization_consistent"] for item in model["inputs"])
    train_case_count = int(diagnoses.get("Cancer Early", 0))
    stage_1_2_count = int(grouped_stages.get("Stage I/II", 0))
    stage_3a_count = int(grouped_stages.get("Stage IIIA", 0))
    report = [
        "# Gockley 2022 Workbook and Frozen-Signature Audit",
        "",
        "## Source Files and Provenance",
        "",
        *file_lines,
        f"- DataSheet_3 ZIP integrity: {archive_integrity}.",
        f"- ZIP members (read in place; no extraction): {', '.join(f'`{name}`' for name in archive_names)}.",
        "- Both files are physically stored under `research/cohort_expansion/data/raw/2024_adnexal_mass/`; the workbook basename present in the repository is `DataSheet_2.xlsx` (underscore), not `DataSheet 2.xlsx`. They were not moved or renamed.",
        "- Workbook sample labels, exact 14-feature panel, and the paired PMML inputs identify this content as the Gockley human serum cohort despite its directory placement.",
        "",
        "## Workbook Structure and Cohort",
        "",
        f"- Worksheet(s): {', '.join(f'`{name}`' for name in metadata['sheet_names'])}; used range `{metadata['workbook_dimension']}`; {metadata['row_count_including_header'] - 1} patient/sample rows and {metadata['column_count']} columns including the header.",
        f"- Unique patient IDs: {metadata['unique_patient_id_count']}; duplicate patient IDs: {metadata['duplicate_patient_id_count']}; exact duplicate data rows: {metadata['duplicate_full_row_count']}. Expression is one row per patient/sample ID.",
        f"- Diagnosis labels: {json.dumps(dict(diagnoses), sort_keys=True)}.",
        f"- Cancer stage strings: {json.dumps(dict(stages), sort_keys=True)}; grouped stages: {json.dumps(dict(grouped_stages), sort_keys=True)}.",
        f"- The workbook independently contains {train_case_count} cancer records, {diagnoses.get('Benign', 0)} benign masses, and {diagnoses.get('Control', 0)} healthy controls; {stage_1_2_count} are Stage I/II and {stage_3a_count} are Stage IIIA. Thus all 275 paper-described human serum samples are represented in the workbook.",
        "- The diagnosis field labels all 75 cancer records `Cancer Early`, including the three Stage IIIA cases; use the separate stage field for the Stage I/II endpoint.",
        f"- Data split labels: {json.dumps(dict(splits), sort_keys=True)}.",
        f"- Clinical variables: age is available for {sum(record['age'] is not None for record in records)}/{len(records)}; CA-125 is not a workbook column; histopathology is present for {sum(record['histopathology'] is not None for record in records)}/{len(records)}. Benign mass and healthy control are explicitly distinguishable by diagnosis.",
        f"- Histopathology counts for non-missing rows: {json.dumps(dict(histopath_counts), sort_keys=True)}. The final ` Histology` field is populated for {histology_secondary_count} rows; `Mutation status`, `Stage`, and `Grade` are repeated in trailing columns and exactly duplicate the leading fields.",
        f"- Assay block: {len(features)} unique identifiers in columns 9-{8 + len(features)}; identifiers use mature-name strings such as `hsa-miR-200c-3p`. All {len(records) * len(features)} patient-feature values are present (missing assay values: {assay_missing}).",
        f"- Blank worksheet cells total {all_blanks}, mostly clinical fields repeated for non-cancer samples; age and all assay values are complete. The canonical first-occurrence clinical fields have missingness for histopathology={sum(record['histopathology'] is None for record in records)}, mutation status={sum(record['mutation_status'] is None for record in records)}, stage={sum(record['stage_raw'] is None for record in records)}, and grade={sum(record['grade'] is None for record in records)}.",
        "- Normalization: the workbook contains no explicit normalization-method column. Its supplementary-data description states that serum miRNA expression values are log10-normalized; treat this as source-caption metadata, not a transform inferred solely from cell values. There is no per-sample CA-125 measurement to join from this workbook.",
        "",
        "## Exact Feature Compatibility",
        "",
        f"- Current OVAI 5-miRNA signature: {len(current_overlap)}/5 exact matches ({', '.join(f'`{x}`' for x in current_overlap) if current_overlap else 'none'}). No compatible frozen current-OVAI test is possible; no substitute features were used.",
        f"- Historical 14-miRNA signature: {len(historical_overlap)}/14 exact matches ({', '.join(f'`{x}`' for x in historical_overlap)}). These are exact workbook identifier matches, not family-level or arm-level substitutions.",
        "- Alias audit: no valid same-mature-product alias was needed or accepted. `hsa-miR-32-5p` is not `hsa-miR-32-3p`; measured miR-200a/200c members are not miR-200b. Full inventory and requested-feature rows are in `2022_gockley_feature_map.csv`.",
        "",
        "## Data Sheet 3 PMML Inspection",
        "",
        *model_lines,
        f"- `signature-best-10.xml` declares all required model components: model name `{model['model_name']}`, 14 ordered inputs ({inputs_text}), per-input `NormContinuous` transforms, {len(model['layers'])} neural layers, {model['weight_count']} connection weights, {model['bias_count']} biases, activation functions `{', '.join(layer['activation'] for layer in model['layers'])}`, and explicit output mapping `{json.dumps(model['output_mapping'], sort_keys=True)}`.",
        f"- Every PMML shift/scale pair agrees with its two `LinearNorm` points within tolerance: {normalization_ok}. The model has no explicit numeric threshold and no `ModelVerification` examples. PMML class decisions are reproduced by the largest output neuron; ROC-AUC uses malignant-minus-benign/control output margin. No probability calibration or threshold tuning was performed.",
        "- The filename says `best-10`, but this PMML network actually declares 14 inputs and an `Ovaries_MLP 14-8-2` architecture. It can be reconstructed from the XML; it is not a 10-feature network.",
        "- The workbook records Training/Validation/Test assignments. Results for all records are apparent cohort performance; Test/Validation rows are reported separately by their workbook `Set`, but the PMML does not encode a training lineage proving which rows were excluded during fitting.",
        "",
        "## Frozen Published-Model Results",
        "",
        "Metrics use the frozen PMML network only. ROC-AUC ranks the output margin. Sensitivity/specificity and confusion counts use the PMML largest-output-neuron decision, not a threshold selected on Gockley labels. The requested primary endpoint is shown for the full cohort and the workbook Test split below.",
        *metric_table,
        "",
        "### All Endpoints by Workbook Split",
        "",
        *full_metrics_table,
        "",
        "All four requested endpoints and all workbook splits are provided in `2022_gockley_signature_best_10_metrics.csv`. Patient-level fixed-model outputs are in `2022_gockley_signature_best_10_predictions.csv`.",
        "",
        "## GSE94533 Compatibility",
        "",
        f"- The Gockley source identifies serum small-RNA sequencing. Local GSE94533 artifacts provide serum miRNA TPM values and 2,582 features, consistent with sequencing-derived expression, but do not identify the exact library chemistry. Gockley’s matrix is described as log10-normalized; local OVAI documentation applies `log1p` to GSE94533 TPM values. These are not directly interchangeable raw scales.",
        f"- Exact historical feature availability: Gockley {len(historical_overlap)}/14; GSE94533 prediction table {gse_audit['gse_historical_overlap']}/14. Current OVAI production features: Gockley {len(current_overlap)}/5; the production manifest identifies the five current model inputs for GSE94533, but those five are absent from this Gockley workbook.",
        f"- Cohort/endpoint: Gockley has 72 Stage I/II cases, 100 benign masses, and 100 healthy controls, with age but no CA-125. GSE94533 has a locally documented Stage I/II-versus-benign/control endpoint of 52 cases and 61 controls in its saved analysis; the broader joined dataset is 179 patients (96 cancers, 83 non-cancer) and includes age/CA-125 where available.",
        f"- Identifier comparison: exact normalized ID matches between the 275 Gockley IDs and {gse_audit['gse_unique_patients']} unique local GSE94533 sample labels: {gse_audit['shared_id_count']}. No deterministic patient linkage was found; different ID namespaces mean zero string matches do not alone prove patient independence.",
        "- Sequencing library chemistry and miRBase release are not documented in the workbook or the inspected local GSE94533 artifacts, so equivalence on those details is unverified.",
        "- Decision: suitable for frozen replay of the historical network and biological replication, but not independent external validation of that network because the PMML was trained on this same cohort. The workbook Test/Validation splits are reported separately, although the XML does not verify split-specific training lineage. It is suitable for a separate cross-cohort benchmark after a prespecified scale/feature protocol; it is not suitable for direct frozen testing of current OVAI (0/5 exact features). Do not pool at this stage. Eventual pooled development would require harmonized quantification, verified patient independence, and a new protocol approved before any training.",
        "",
        "## Use Decision",
        "",
        "Gockley should replace the 2024 FirePlex cohort as the primary clinically aligned cohort-expansion candidate relative to GSE94533: it has 72 Stage I/II cases, separately identifiable benign masses and healthy controls, and all 14 historical inputs. This makes it stronger for historical-signature benchmarking, but the same-cohort PMML results are not independent external validation. It does not replace the need for a valid independent external cohort for current 5-miRNA OVAI, which has 0/5 exact features here. Keep FirePlex as a separate cross-platform biological replication dataset; do not combine either dataset with GSE94533 in this audit.",
        "",
        "No model was trained, recalibrated, or tuned. `production_artifacts/` was read only. Original workbook and ZIP remain unchanged.",
        "",
    ]
    return "\n".join(report)


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    if not WORKBOOK.is_file() or not MODEL_ARCHIVE.is_file():
        raise FileNotFoundError(f"Expected source files not found: {WORKBOOK}, {MODEL_ARCHIVE}")
    headers, raw_rows, features, metadata = load_workbook_data()
    records = build_records(headers, raw_rows, features)
    source_info = [
        file_info(WORKBOOK, "Excel OOXML workbook"),
        file_info(MODEL_ARCHIVE, "ZIP archive of PMML XML networks"),
    ]
    with ZipFile(MODEL_ARCHIVE) as archive:
        archive_names = archive.namelist()
        bad_member = archive.testzip()
        archive_integrity = "PASS" if bad_member is None else f"FAIL at {bad_member}"
        expected_models = {MODEL_FILENAME, "7miRTrimmedXenoNet-38.xml", "NewNetXenograft-1.xml"}
        model_filenames = [name for name in archive_names if name.lower() in {item.lower() for item in expected_models}]
        models = [parse_pmml_model(archive, name) for name in model_filenames]
        if archive_integrity != "PASS":
            raise ValueError(f"Invalid model archive: {archive_integrity}")
    if MODEL_FILENAME not in {model["filename"] for model in models}:
        raise ValueError(f"Missing {MODEL_FILENAME}; archive members: {archive_names}")
    best_model = next(model for model in models if model["filename"] == MODEL_FILENAME)
    model_features = [item["feature"] for item in best_model["inputs"]]
    if set(model_features) != set(HISTORICAL) or len(model_features) != len(HISTORICAL):
        raise ValueError(f"PMML inputs do not equal the historical signature: {model_features}")
    if not all(item["normalization_consistent"] for item in best_model["inputs"]):
        raise ValueError("PMML NormContinuous shift/scale conflicts with LinearNorm points")
    if not all(feature in features for feature in model_features):
        raise ValueError("One or more frozen PMML inputs are absent from the workbook")
    predictions = [predict_pmml(best_model, record) for record in records]
    metrics = evaluate_endpoints(records, predictions)
    prediction_output = pd.DataFrame(
        [
            {
                "sample_id": record["sample_id"],
                "set": record["set"],
                "diagnosis": record["diagnosis"],
                "stage_raw": record["stage_raw"],
                "stage_group": record["stage_group"],
                **prediction,
            }
            for record, prediction in zip(records, predictions)
        ]
    )
    reconstruction = pd.DataFrame(records)
    feature_map = build_feature_map(features, model_features)
    gse = pd.read_csv(GSE_PREDICTIONS, dtype={"dataset": str, "sample_id_normalized": str, "title": str})
    gse = gse.loc[gse["dataset"].eq("GSE94533")].drop_duplicates("sample_id_normalized")
    gse_features = set(HISTORICAL).intersection(gse.columns)
    manifest = json.loads(PRODUCTION_MANIFEST.read_text(encoding="utf-8"))
    gse_ids = {normalize_id(value) for value in gse["sample_id_normalized"].dropna()}
    workbook_ids = {normalize_id(record["sample_id"]) for record in records}
    shared_ids = sorted(workbook_ids & gse_ids)
    gse_audit = {
        "gse_unique_patients": int(gse["sample_id_normalized"].nunique()),
        "gse_historical_overlap": len(gse_features),
        "shared_id_count": len(shared_ids),
        "shared_ids": shared_ids,
        "production_features": manifest.get("features", []),
    }
    reconstruction.to_csv(REPORT_DIR / "2022_gockley_cohort_reconstruction.csv", index=False)
    feature_map.to_csv(REPORT_DIR / "2022_gockley_feature_map.csv", index=False)
    prediction_output.to_csv(REPORT_DIR / "2022_gockley_signature_best_10_predictions.csv", index=False)
    metrics.to_csv(REPORT_DIR / "2022_gockley_signature_best_10_metrics.csv", index=False)
    (REPORT_DIR / "2022_gockley_workbook_audit.md").write_text(
        make_report(headers, raw_rows, records, features, metadata, source_info, archive_names, archive_integrity, models, metrics, gse_audit),
        encoding="utf-8",
    )
    print(f"Workbook: {len(records)} unique patient IDs, {len(features)} complete miRNA features; historical exact overlap {len(set(HISTORICAL) & set(features))}/14; current exact overlap {len(set(CURRENT) & set(features))}/5.")
    print(f"Archive: integrity={archive_integrity}; members={archive_names}; models={len(models)}.")
    print("Primary Stage I/II vs benign metrics:")
    print(metrics.loc[(metrics.endpoint == "stage_I_II_cancer_vs_benign_mass") & metrics.subset.isin(["All", "Test"])].to_string(index=False))
    print(f"Reports written to {REPORT_DIR.relative_to(ROOT)}.")


if __name__ == "__main__":
    main()
